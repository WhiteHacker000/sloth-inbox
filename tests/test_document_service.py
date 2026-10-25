from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from backend.document_service import (
    DocumentError,
    extract_document_text,
    read_text_bytes,
    read_text_file,
    validate_text,
)


def make_text_pdf(text: str | None) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text is not None:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_ref}
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(
            b"BT /F1 12 Tf 72 240 Td (" + text.encode("ascii") + b") Tj ET"
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Project update")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Rahul"
    document.add_paragraph("Next steps")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_accepts_valid_text_and_utf8_bom() -> None:
    assert validate_text("A document", 100) == "A document"
    assert read_text_bytes(b"\xef\xbb\xbfA document", 100) == "A document"


def test_rejects_empty_input() -> None:
    with pytest.raises(DocumentError, match="empty"):
        validate_text(" \n\t", 100)


def test_rejects_invalid_utf8() -> None:
    with pytest.raises(DocumentError, match="UTF-8"):
        read_text_bytes(b"\xff\xfe\x81", 100)


def test_rejects_oversized_text_and_upload() -> None:
    with pytest.raises(DocumentError, match="size limit"):
        validate_text("too long", 3)
    with pytest.raises(DocumentError, match="size limit"):
        read_text_bytes(b"too long", 3)


def test_reads_text_file(tmp_path: Path) -> None:
    document = tmp_path / "brief.txt"
    document.write_text("Project update", encoding="utf-8")
    assert read_text_file(document, 100) == "Project update"


def test_missing_file_and_unsupported_format_are_rejected(tmp_path: Path) -> None:
    with pytest.raises(DocumentError, match="does not exist"):
        read_text_file(tmp_path / "missing.txt", 100)

    pdf = tmp_path / "document.pdf"
    pdf.write_text("not a PDF", encoding="utf-8")
    with pytest.raises(DocumentError, match="Only plain-text"):
        read_text_file(pdf, 100)


def test_extracts_unicode_txt_upload() -> None:
    assert extract_document_text("Résumé – 更新".encode(), "notes.txt", 100) == (
        "Résumé – 更新"
    )


def test_extracts_selectable_pdf_text() -> None:
    assert extract_document_text(make_text_pdf("Project update"), "brief.pdf", 1000) == (
        "Project update"
    )


def test_reports_pdf_without_selectable_text_as_ocr_unsupported() -> None:
    with pytest.raises(DocumentError, match="Scanned PDFs requiring OCR"):
        extract_document_text(make_text_pdf(None), "scan.pdf", 1000)


def test_rejects_malformed_pdf() -> None:
    with pytest.raises(DocumentError, match="malformed|valid PDF signature"):
        extract_document_text(b"%PDF-1.4\nnot a real PDF", "broken.pdf", 1000)


def test_extracts_docx_paragraphs_and_table_in_document_order() -> None:
    assert extract_document_text(make_docx(), "brief.docx", 100000) == (
        "Project update\nOwner | Rahul\nNext steps"
    )


def test_rejects_docx_without_readable_text() -> None:
    document = Document()
    output = BytesIO()
    document.save(output)

    with pytest.raises(DocumentError, match="empty"):
        extract_document_text(output.getvalue(), "empty.docx", 100000)


def test_rejects_malformed_docx() -> None:
    with pytest.raises(DocumentError, match="DOCX"):
        extract_document_text(b"PK\x03\x04invalid archive", "broken.docx", 1000)


def test_rejects_empty_and_unsupported_uploads() -> None:
    with pytest.raises(DocumentError, match="empty"):
        extract_document_text(b"", "empty.txt", 100)
    with pytest.raises(DocumentError, match="Unsupported document type"):
        extract_document_text(b"content", "notes.rtf", 100)


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("large.txt", b"long"),
        ("large.pdf", b"%PDF-1.4" + b"x" * 10),
        ("large.docx", b"PK" + b"x" * 10),
    ],
)
def test_rejects_oversized_upload_before_parsing(
    filename: str,
    content: bytes,
) -> None:
    max_size_bytes = 3 if filename.endswith(".txt") else 5
    with pytest.raises(DocumentError, match="size limit"):
        extract_document_text(content, filename, max_size_bytes)


def test_rejects_wrong_file_signature_for_binary_formats() -> None:
    with pytest.raises(DocumentError, match="valid PDF signature"):
        extract_document_text(b"not PDF text", "fake.pdf", 100)

    with pytest.raises(DocumentError, match="valid DOCX"):
        extract_document_text(b"not a zip file", "fake.docx", 100)
